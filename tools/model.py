# -*- coding: utf-8 -*-
"""Datový model a parser studijních materiálů ISTQB CTFL.

Čte .docx soubory ve složce projektu a převádí je na strukturu bloků.
Dokumenty nemají styly nadpisů — struktura je nesena velikostí písma:

    24 pt tučně  -> nadpis dokumentu (H1)
    13 pt šedě   -> podtitul dokumentu
    17 pt tučně  -> nadpis oddílu (H2)
    ostatní      -> odstavec

Barva písma nese význam, ale v každé variantě jiný:

    výklad / zkrácená    0B5394 = doplnění nad rámec sylabu
                         1B6E3C = příklad z praxe
    otázky a odpovědi    0B5394 = K1, C0392B = K2, 1B6E3C = K3
                         tučně  = otázka, netučně = odpověď

Invariant: text se nikde nemění ani nevypouští.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import docx

# --- barvy z dokumentů ------------------------------------------------------

MODRA = "0B5394"
ZELENA = "1B6E3C"
CERVENA = "C0392B"
TEXT = "1A1A1A"
SEDA = "555555"

# barva -> kognitivní úroveň (platí jen ve variantě "otazky")
K_UROVEN = {MODRA: "K1", CERVENA: "K2", ZELENA: "K3"}

# barva -> druh vsuvky (platí ve variantě "vyklad" a "zkracena")
VSUVKA = {MODRA: "nad-ramec", ZELENA: "praxe"}

VSUVKA_POPIS = {
    "nad-ramec": "Nad rámec sylabu",
    "praxe": "Příklad z praxe",
}

# odstavce, které uzavírají dokument a nepatří k předchozí odpovědi
UZAVIRACI_PREFIXY = ("Tím končí",)

VARIANTY = {
    "PLNY-VYKLAD": "vyklad",
    "ZKRACENA-VERZE": "zkracena",
    "OTAZKY-A-ODPOVEDI": "otazky",
}

VARIANTA_NAZEV = {
    "vyklad": "Plný výklad",
    "zkracena": "Zkrácená verze",
    "otazky": "Otázky a odpovědi",
}

VARIANTA_POPIS = {
    "vyklad": "Souvislý výklad celé kapitoly s doplňky nad rámec sylabu a příklady z praxe.",
    "zkracena": "Jen to, co je nutné znát ke zkoušce. Bez odboček, k rychlému opakování.",
    "otazky": "Otázky ke všem studijním cílům, barevně odlišené podle kognitivní úrovně.",
}

# minuty výuky podle sylabu CTFL v4.0 (celkem 1 135)
MINUTY_VYUKY = {1: 180, 2: 130, 3: 80, 4: 390, 5: 335, 6: 20}

KAPITOLY = {
    1: "Základy testování",
    2: "Testování v průběhu životního cyklu vývoje softwaru",
    3: "Statické testování",
    4: "Testovací analýza a návrh testů",
    5: "Management testování",
    6: "Testovací nástroje",
}

KAPITOLY_KRATCE = {
    1: "Základy testování",
    2: "Testování v průběhu životního cyklu",
    3: "Statické testování",
    4: "Testovací analýza a návrh testů",
    5: "Management testování",
    6: "Testovací nástroje",
}


# --- pomocné ----------------------------------------------------------------


def bez_diakritiky(text: str) -> str:
    rozlozeno = unicodedata.normalize("NFKD", text)
    return "".join(z for z in rozlozeno if not unicodedata.combining(z))


def slug(text: str) -> str:
    zaklad = bez_diakritiky(text).lower()
    zaklad = re.sub(r"[^a-z0-9]+", "-", zaklad).strip("-")
    return zaklad or "oddil"


def pocet_slov(text: str) -> int:
    return len(text.split())


# --- bloky ------------------------------------------------------------------


@dataclass
class Nadpis:
    text: str
    kotva: str


@dataclass
class Odstavec:
    text: str


@dataclass
class Vsuvka:
    druh: str  # "nad-ramec" | "praxe"
    odstavce: list[str] = field(default_factory=list)


@dataclass
class Legenda:
    uroven: str  # "K1" | "K2" | "K3"
    text: str


@dataclass
class Otazka:
    uroven: str
    cislo: int | None
    kotva: str
    text: str
    odpovedi: list[str] = field(default_factory=list)


@dataclass
class Dokument:
    kapitola: int
    varianta: str
    nadpis: str
    podtitul: str
    bloky: list = field(default_factory=list)
    zdroj: Path | None = None

    @property
    def oddily(self) -> list[Nadpis]:
        return [b for b in self.bloky if isinstance(b, Nadpis)]

    @property
    def pocet_otazek(self) -> int:
        return sum(1 for b in self.bloky if isinstance(b, Otazka))

    def slova_zdroje(self) -> int:
        """Počet slov ve zdrojovém .docx — referenční hodnota pro kontrolu."""
        d = docx.Document(str(self.zdroj))
        return sum(pocet_slov(p.text.strip()) for p in d.paragraphs if p.text.strip())


# --- čtení docx -------------------------------------------------------------


def _vlastnosti(p) -> tuple[float | None, bool, str | None]:
    """Velikost písma, tučnost a barva prvního běhu odstavce."""
    beh = p.runs[0] if p.runs else None
    if beh is None:
        return None, False, None
    velikost = beh.font.size.pt if beh.font.size else None
    tucne = bool(beh.bold)
    barva = None
    try:
        if beh.font.color is not None and beh.font.color.type is not None:
            rgb = beh.font.color.rgb
            if rgb is not None:
                barva = str(rgb).upper()
    except (AttributeError, ValueError):
        barva = None
    return velikost, tucne, barva


def _je_uzaviraci(text: str) -> bool:
    return any(text.startswith(p) for p in UZAVIRACI_PREFIXY)


def nacti(cesta: Path) -> Dokument:
    nazev = cesta.stem
    kapitola = int(re.search(r"_Kap(\d)_", nazev).group(1))
    pripona = nazev.rsplit("_", 1)[-1]
    varianta = VARIANTY[pripona]

    doc = docx.Document(str(cesta))
    vysledek = Dokument(kapitola=kapitola, varianta=varianta, nadpis="", podtitul="", zdroj=cesta)

    pouzite_kotvy: set[str] = set()

    def unikatni(zaklad: str) -> str:
        kotva = zaklad
        n = 2
        while kotva in pouzite_kotvy:
            kotva = f"{zaklad}-{n}"
            n += 1
        pouzite_kotvy.add(kotva)
        return kotva

    otevrena_vsuvka: Vsuvka | None = None
    otevrena_otazka: Otazka | None = None

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        velikost, tucne, barva = _vlastnosti(p)

        if velikost and velikost >= 24:
            vysledek.nadpis = text
            continue
        if velikost and 12 <= velikost <= 14:
            vysledek.podtitul = text
            continue
        if velikost and velikost >= 17:
            otevrena_vsuvka = None
            otevrena_otazka = None
            vysledek.bloky.append(Nadpis(text=text, kotva=unikatni(slug(text))))
            continue

        if varianta == "otazky":
            uroven = K_UROVEN.get(barva or "")
            if uroven and tucne:
                shoda = re.match(r"^(\d+)\.\s", text)
                cislo = int(shoda.group(1)) if shoda else None
                kotva = unikatni(f"otazka-{cislo}" if cislo else slug(text[:40]))
                otevrena_otazka = Otazka(uroven=uroven, cislo=cislo, kotva=kotva, text=text)
                vysledek.bloky.append(otevrena_otazka)
            elif uroven:
                otevrena_otazka = None
                vysledek.bloky.append(Legenda(uroven=uroven, text=text))
            elif otevrena_otazka is not None and not _je_uzaviraci(text):
                otevrena_otazka.odpovedi.append(text)
            else:
                otevrena_otazka = None
                vysledek.bloky.append(Odstavec(text=text))
            continue

        druh = VSUVKA.get(barva or "")
        if druh:
            if otevrena_vsuvka is not None and otevrena_vsuvka.druh == druh:
                otevrena_vsuvka.odstavce.append(text)
            else:
                otevrena_vsuvka = Vsuvka(druh=druh, odstavce=[text])
                vysledek.bloky.append(otevrena_vsuvka)
        else:
            otevrena_vsuvka = None
            vysledek.bloky.append(Odstavec(text=text))

    return vysledek


def nacti_vse(slozka: Path) -> dict[tuple[int, str], Dokument]:
    dokumenty: dict[tuple[int, str], Dokument] = {}
    for cesta in sorted(slozka.glob("CTFL_Kap*.docx")):
        if cesta.name.startswith("~$"):
            continue
        d = nacti(cesta)
        dokumenty[(d.kapitola, d.varianta)] = d
    return dokumenty
